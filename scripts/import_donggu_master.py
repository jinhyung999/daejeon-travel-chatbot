# =====================================================
# import_donggu_master.py
# "00. 동구 업소 마스터 데이터.xlsx" -> place 테이블 적재
#
# 엑셀 특이사항:
#   - 주소가 지번주소만 있고 도로명주소/위도경도가 없음
#     -> 카카오 로컬 주소검색 API로 지번 -> 도로명 + 좌표 동시 변환
#   - "업소 분류"(tag)가 자유 텍스트라 place.category enum으로 매핑 필요
#
# 중복 판정: scripts/dedupe_place.py와 동일한 기준(정규화 이름 동일
#            + 50m 이내)을 기존 place 테이블 대상으로 적용, 겹치면 skip
#
# 사용법:
#   python scripts/import_donggu_master.py            # dry-run (집계만 출력)
#   python scripts/import_donggu_master.py --apply     # 실제 INSERT
# =====================================================

import json
import math
import os
import re
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl
import requests
from dotenv import load_dotenv

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from collectors.common import get_conn  # noqa: E402

load_dotenv()

XLSX_PATH = REPO_ROOT / "data" / "00. 동구 업소 마스터 데이터.xlsx"
GEOCODE_CACHE_PATH = REPO_ROOT / "data" / "raw" / "donggu_geocode_cache.json"
KAKAO_ADDRESS_URL = "https://dapi.kakao.com/v2/local/search/address.json"
KAKAO_KEY = os.getenv("KAKAO_REST_API_KEY")

SOURCE_API = "donggu_sbiz_master"
DIST_THRESHOLD_M = 50  # scripts/dedupe_place.py와 동일 기준

# 업소 분류(tag) -> place.category 매핑 (사용자 확정)
CATEGORY_MAP = {
    "음식점": "restaurant",
    "카페": "cafe",
    "북카페": "cafe",
    "서점": "culture",
    "도서관": "culture",
    "갤러리": "culture",
    "박물관": "culture",
    "전시관": "culture",
    "미술관": "culture",
    "복합문화공간": "culture",
    "공연장": "culture",
    "연극시설": "culture",
    "영화관": "culture",
    "문화재": "attraction",
    "공원": "attraction",
    "특화거리": "attraction",
    "시장": "shopping",
}
# 여행과 무관 -> 제외 (커튼/한복/한약방/건어물/건축자재/공구/연관없음)

HEADER_ROW = 2  # 1-based: 2번째 행이 영문 키
DATA_START_ROW = 3

COLS = [
    "total_id", "naver_id", "google_id", "title", "jibun_address",
    "admin_dong", "address_dong", "phone_number_naver", "phone_number_google",
    "business_hours", "peak_time", "peak_day_1st", "peak_day_2nd",
    "peak_time_1st", "peak_time_2nd", "menu", "age_based_popularity",
    "gender_male_popularity", "gender_female_popularity", "services_provided",
    "services_not_provided", "price_range", "theme_keyword",
    "keyword_reviews", "review_score_naver", "review_score_google",
    "review_all_count_google", "broadcasting_info", "auth_info", "tag",
]

EXTRA_JSON_FIELDS = [
    "naver_id", "google_id", "jibun_address", "admin_dong", "address_dong",
    "peak_time", "peak_day_1st", "peak_day_2nd", "peak_time_1st",
    "peak_time_2nd", "age_based_popularity", "gender_male_popularity",
    "gender_female_popularity", "services_provided", "services_not_provided",
    "price_range", "keyword_reviews", "review_score_naver",
    "review_score_google", "review_all_count_google", "broadcasting_info",
    "auth_info",
]
JSON_TEXT_FIELDS = {"services_provided", "services_not_provided", "keyword_reviews"}


def haversine(lat1, lng1, lat2, lng2):
    r = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def norm_name(name):
    return (name or "").replace(" ", "")


def load_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["Result 1"]
    rows = []
    for values in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if values[0] is None:
            continue
        rows.append(dict(zip(COLS, values)))
    return rows


def load_geocode_cache():
    if GEOCODE_CACHE_PATH.exists():
        with open(GEOCODE_CACHE_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_geocode_cache(cache):
    GEOCODE_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(GEOCODE_CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)


# 시/도 + (구) + 동/리/가 + (산)번지 핵심부만 뽑아내는 패턴.
# 엑셀 지번주소에 "동구" 표기가 빠졌거나 번지 뒤에 건물명이 붙어있어
# 카카오 API가 못 찾는 경우, 이 핵심부만 남겨 재시도한다.
JIBUN_CORE_RE = re.compile(
    r"(?P<city>\S+광역시|\S+특별시|\S+도)\s*(?P<gu>\S+구)?\s*"
    r"(?P<dong>\S+(?:동|리|가))\s+(?P<lot>산\s?\d+(?:-\d+)?|\d+(?:-\d+)?)"
)


def _trim_jibun(jibun_address):
    m = JIBUN_CORE_RE.search(jibun_address)
    if not m:
        return None
    gu = m.group("gu") or "동구"
    trimmed = f"{m.group('city')} {gu} {m.group('dong')} {m.group('lot')}".replace("산 ", "산")
    return trimmed if trimmed != jibun_address else None


def _geocode_once(query):
    for attempt in range(3):
        resp = requests.get(
            KAKAO_ADDRESS_URL,
            headers={"Authorization": f"KakaoAK {KAKAO_KEY}"},
            params={"query": query},
            timeout=10,
        )
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        resp.raise_for_status()
        docs = resp.json().get("documents") or []
        if docs:
            doc = docs[0]
            road = doc.get("road_address")
            if road and road.get("address_name"):
                return {
                    "road_address": road["address_name"],
                    "lat": float(doc["y"]),
                    "lng": float(doc["x"]),
                }
        return None
    return None


def geocode(jibun_address, cache):
    if jibun_address in cache:
        return cache[jibun_address]

    result = _geocode_once(jibun_address)
    if result is None:
        trimmed = _trim_jibun(jibun_address)
        if trimmed:
            result = _geocode_once(trimmed)

    cache[jibun_address] = result
    return result


def build_extra_json(row):
    extra = {}
    for field in EXTRA_JSON_FIELDS:
        value = row.get(field)
        if value in (None, ""):
            continue
        if field in JSON_TEXT_FIELDS and isinstance(value, str):
            try:
                value = json.loads(value)
            except (json.JSONDecodeError, TypeError):
                pass
        extra[field] = value
    return json.dumps(extra, ensure_ascii=False) if extra else None


def detect_has_parking(row):
    text = row.get("services_provided") or ""
    if not isinstance(text, str):
        text = json.dumps(text, ensure_ascii=False)
    return 1 if "주차" in text else None


def load_existing_index(conn, categories):
    placeholders = ",".join("?" for _ in categories)
    cur = conn.execute(
        f"SELECT name, address, lat, lng FROM place "
        f"WHERE category IN ({placeholders}) AND lat IS NOT NULL AND lng IS NOT NULL",
        categories,
    )
    grid = defaultdict(list)
    for name, address, lat, lng in cur.fetchall():
        key = (round(lat, 3), round(lng, 3))
        grid[key].append((norm_name(name), address, lat, lng))
    return grid


def is_duplicate(grid, name, address, lat, lng):
    n = norm_name(name)
    base_key = (round(lat, 3), round(lng, 3))
    for dlat in (-1, 0, 1):
        for dlng in (-1, 0, 1):
            key = (round(base_key[0] + dlat * 0.001, 3), round(base_key[1] + dlng * 0.001, 3))
            for ex_name, ex_addr, ex_lat, ex_lng in grid.get(key, []):
                if ex_name != n:
                    continue
                if (address and ex_addr and address == ex_addr) or \
                        haversine(lat, lng, ex_lat, ex_lng) <= DIST_THRESHOLD_M:
                    return True
    return False


def main():
    apply = "--apply" in sys.argv
    limit = None
    for arg in sys.argv:
        if arg.startswith("--limit="):
            limit = int(arg.split("=", 1)[1])

    if not KAKAO_KEY or KAKAO_KEY == "your_kakao_rest_api_key_here":
        print("KAKAO_REST_API_KEY가 설정되어 있지 않습니다.")
        sys.exit(1)

    rows = load_rows()
    if limit:
        rows = rows[:limit]
    print(f"엑셀 원본: {len(rows)}건")

    cache = load_geocode_cache()

    stats = defaultdict(int)
    prepared = []

    conn = get_conn()
    grid = load_existing_index(conn, sorted(set(CATEGORY_MAP.values())))

    skipped_tags = defaultdict(int)
    geocode_fail = []

    for i, row in enumerate(rows, 1):
        tag = row.get("tag")
        category = CATEGORY_MAP.get(tag)
        if category is None:
            skipped_tags[tag] += 1
            stats["excluded_category"] += 1
            continue

        jibun = row.get("jibun_address")
        if not jibun:
            stats["no_address"] += 1
            continue

        geo = geocode(jibun, cache)
        if geo is None:
            geocode_fail.append((row.get("title"), jibun))
            stats["geocode_fail"] += 1
            continue

        name = row.get("title")
        if is_duplicate(grid, name, geo["road_address"], geo["lat"], geo["lng"]):
            stats["duplicate"] += 1
            continue

        tel = row.get("phone_number_naver") or row.get("phone_number_google")
        prepared.append({
            "place_id": f"donggu_sbiz_{row['total_id']}",
            "name": name,
            "category": category,
            "address": geo["road_address"],
            "lat": geo["lat"],
            "lng": geo["lng"],
            "open_time": row.get("business_hours"),
            "close_day": None,
            "fee": None,
            "has_parking": detect_has_parking(row),
            "tel": tel,
            "source_api": SOURCE_API,
            "extra_json": build_extra_json(row),
            "overview": None,
            "homepage": None,
            "recommend": None,
            "concept_tag": row.get("theme_keyword"),
            "photo_spot": None,
            "has_workshop": None,
            "signature_menu": row.get("menu"),
        })
        stats["inserted_candidate"] += 1
        if prepared[-1]["has_parking"] == 1:
            stats["has_parking_detected"] += 1
        # 새로 넣는 것도 같은 배치 내 중복 방지를 위해 인덱스에 반영
        key = (round(geo["lat"], 3), round(geo["lng"], 3))
        grid[key].append((norm_name(name), geo["road_address"], geo["lat"], geo["lng"]))

        if i % 200 == 0:
            print(f"  진행 {i}/{len(rows)} (geocode 캐시 {len(cache)}건)")
            save_geocode_cache(cache)

    save_geocode_cache(cache)

    print("\n=== 집계 ===")
    print(f"제외(분류 무관): {stats['excluded_category']}건")
    for tag, cnt in sorted(skipped_tags.items(), key=lambda x: -x[1]):
        print(f"    - {tag}: {cnt}건")
    print(f"주소 없음: {stats['no_address']}건")
    print(f"지오코딩 실패: {stats['geocode_fail']}건")
    print(f"기존 place와 중복(skip): {stats['duplicate']}건")
    print(f"신규 삽입 대상: {stats['inserted_candidate']}건")
    print(f"  - 그중 has_parking=1 감지: {stats['has_parking_detected']}건")

    if geocode_fail:
        fail_path = REPO_ROOT / "data" / "raw" / "donggu_geocode_fail.json"
        with open(fail_path, "w", encoding="utf-8") as f:
            json.dump(geocode_fail, f, ensure_ascii=False, indent=2)
        print(f"지오코딩 실패 목록 저장: {fail_path}")

    if not apply:
        print("\n[dry-run] 실제 반영하려면: python scripts/import_donggu_master.py --apply")
        conn.close()
        return

    cur = conn.cursor()
    cur.executemany("""
        INSERT INTO place (place_id, name, category, address, lat, lng,
                            open_time, close_day, fee, has_parking, tel,
                            source_api, extra_json, overview, homepage,
                            recommend, concept_tag, photo_spot, has_workshop,
                            signature_menu)
        VALUES (:place_id, :name, :category, :address, :lat, :lng,
                :open_time, :close_day, :fee, :has_parking, :tel,
                :source_api, :extra_json, :overview, :homepage,
                :recommend, :concept_tag, :photo_spot, :has_workshop,
                :signature_menu)
    """, prepared)
    conn.commit()
    total = cur.execute("SELECT COUNT(*) FROM place").fetchone()[0]
    print(f"\n삽입 완료: {len(prepared)}건 (place 테이블 총 {total}행)")
    conn.close()


if __name__ == "__main__":
    main()
